"""Handelski eksport sklasyfikowanych dostaw B2B."""

from __future__ import annotations

import csv
from datetime import date
from io import BytesIO, StringIO
from zipfile import is_zipfile

import pytest
from openpyxl import load_workbook

from core.delivery_memory import ClassifiedDeliveryItem
from core.export_package import write_csv, write_xlsx
from core.models import DeliveryCategory, Investment


def make_item(
    case_id: str,
    category: DeliveryCategory,
    **investment_fields: object,
) -> ClassifiedDeliveryItem:
    investment = Investment(
        id_sprawy=case_id,
        zrodlo="pozwolenia",
        status="decyzja",
        **investment_fields,
    )
    return ClassifiedDeliveryItem(investment, f"revision-{case_id}", category)


def test_xlsx_has_three_delivery_sheets_with_scoring_timing_and_location() -> None:
    items = (
        make_item(
            "CASE-NEW",
            DeliveryCategory.NEW,
            organ="Starostwo Powiatowe",
            data_decyzji="2025-09-01",
            kategoria="mieszkaniowa-jednorodzinna",
            nazwa_zamierzenia="Budowa zespołu dwóch budynków mieszkalnych",
            miejscowosc="Łódź",
            powiat="łódzki",
            gmina="Łódź",
            dzialki=["123/4"],
            geoportal_url="https://example.test/geoportal",
            google_maps_url="https://example.test/maps",
        ),
        make_item("CASE-CHANGE", DeliveryCategory.CHANGE),
        make_item("CASE-FULL", DeliveryCategory.FULL),
    )
    output = BytesIO()

    write_xlsx(items, output, trade_key="dach", as_of=date(2026, 1, 1))

    payload = output.getvalue()
    assert is_zipfile(BytesIO(payload))
    workbook = load_workbook(BytesIO(payload), data_only=False)
    assert workbook.sheetnames == ["1. Nowe", "2. Zmiany", "3. Pełna baza"]

    new_sheet = workbook["1. Nowe"]
    header_columns = {cell.value: cell.column for cell in new_sheet[1]}
    assert new_sheet.max_row == 2
    assert new_sheet.cell(2, header_columns["ID sprawy"]).value == "CASE-NEW"
    assert new_sheet.cell(2, header_columns["Organ"]).value == "Starostwo Powiatowe"
    assert new_sheet.cell(2, header_columns["Miejscowość"]).value == "Łódź"
    assert new_sheet.cell(2, header_columns["Działki"]).value == "123/4"
    assert new_sheet.cell(2, header_columns["Uzasadnienie"]).value
    assert "Etap dachu" in new_sheet.cell(2, header_columns["Kiedy dzwonić"]).value
    assert new_sheet.cell(2, header_columns["Geoportal"]).value == "https://example.test/geoportal"
    assert new_sheet.cell(2, header_columns["Google Maps"]).value == "https://example.test/maps"
    assert workbook["2. Zmiany"].max_row == 2
    assert workbook["3. Pełna baza"].max_row == 2
    workbook.close()


def test_csv_is_flat_utf8_bom_export_with_polish_text_and_csv_quoting() -> None:
    item = make_item(
        "CASE-ŁÓDŹ",
        DeliveryCategory.NEW,
        nazwa_zamierzenia='Budowa "Domu", w Łodzi',
        miejscowosc="Łódź",
    )
    output = StringIO()

    write_csv((item,), output)

    value = output.getvalue()
    assert value.startswith("\ufeff")
    rows = list(csv.reader(StringIO(value.removeprefix("\ufeff"))))
    header_columns = {name: index for index, name in enumerate(rows[0])}
    assert len(rows) == 2
    assert rows[1][header_columns["ID sprawy"]] == "CASE-ŁÓDŹ"
    assert rows[1][header_columns["Nazwa zamierzenia"]] == 'Budowa "Domu", w Łodzi'
    assert rows[1][header_columns["Miejscowość"]] == "Łódź"


def test_spreadsheet_formula_text_is_kept_as_literal_in_xlsx_and_escaped_in_csv() -> None:
    formula_text = "=HYPERLINK(\"https://example.test\",\"lead\")"
    item = make_item("CASE-FORMULA", DeliveryCategory.NEW, inwestor=formula_text)
    xlsx_output = BytesIO()

    write_xlsx((item,), xlsx_output)

    workbook = load_workbook(BytesIO(xlsx_output.getvalue()), data_only=False)
    sheet = workbook["1. Nowe"]
    header_columns = {cell.value: cell.column for cell in sheet[1]}
    investor_cell = sheet.cell(2, header_columns["Inwestor"])
    assert investor_cell.value == formula_text
    assert investor_cell.data_type != "f"
    workbook.close()

    csv_output = StringIO()
    write_csv((item,), csv_output)
    rows = list(csv.reader(StringIO(csv_output.getvalue().removeprefix("\ufeff"))))
    csv_headers = {name: index for index, name in enumerate(rows[0])}
    assert rows[1][csv_headers["Inwestor"]] == "'" + formula_text


@pytest.mark.parametrize("formula_text", ("=1+1", "+1+1", "-1+1", "@SUM(1,1)", " \t=1+1"))
def test_csv_escapes_each_supported_formula_prefix(formula_text: str) -> None:
    item = make_item("CASE-FORMULA", DeliveryCategory.NEW, inwestor=formula_text)
    output = StringIO()

    write_csv((item,), output)

    rows = list(csv.reader(StringIO(output.getvalue().removeprefix("\ufeff"))))
    investor_column = rows[0].index("Inwestor")
    assert rows[1][investor_column] == "'" + formula_text


def test_xlsx_requires_reference_date_when_a_trade_window_is_requested() -> None:
    with pytest.raises(ValueError, match="as_of"):
        write_xlsx((make_item("CASE-1", DeliveryCategory.NEW),), BytesIO(), trade_key="dach")
